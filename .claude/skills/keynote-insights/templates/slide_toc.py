"""Table of contents — two-column list of selected features."""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from . import theme as T


def render_toc(prs, *, title: str, items: list[tuple[str, int]]):
    """items = [(feature_title_zh, page_number), ...]"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Light gray background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, T.SLIDE_W, T.SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = T.LIGHT_GRAY; bg.line.fill.background()

    # Left crimson vertical bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(0.6), Inches(0.8),
                                 Inches(0.18), Inches(6.0))
    bar.fill.solid(); bar.fill.fore_color.rgb = T.CRIMSON; bar.line.fill.background()

    # Title
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(0.7), Inches(11), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    T.apply_text(p.add_run(), title,
                 font_name=T.FONT_CN, size=T.FONT_SIZE["closing_title"],
                 bold=True, color=T.CRIMSON)

    # Items: 2 columns
    mid = (len(items) + 1) // 2
    cols = [items[:mid], items[mid:]]
    col_lefts = [1.0, 7.2]
    for c, col in enumerate(cols):
        for r, (name, page) in enumerate(col):
            top = 2.0 + r * 0.55
            tb = slide.shapes.add_textbox(Inches(col_lefts[c]), Inches(top),
                                          Inches(5.5), Inches(0.5))
            tp = tb.text_frame.paragraphs[0]
            T.apply_text(tp.add_run(), f"{c*mid + r + 1:02d}  {name}",
                         font_name=T.FONT_CN, size=T.FONT_SIZE["toc_item"],
                         color=T.BODY_GRAY)
            pg = slide.shapes.add_textbox(Inches(col_lefts[c] + 5.0), Inches(top),
                                          Inches(0.8), Inches(0.5))
            pp = pg.text_frame.paragraphs[0]
            pp.alignment = PP_ALIGN.RIGHT
            T.apply_text(pp.add_run(), f"{page:02d}",
                         font_name=T.FONT_EN, size=T.FONT_SIZE["toc_item"],
                         color=T.CRIMSON, bold=True)
    return slide
