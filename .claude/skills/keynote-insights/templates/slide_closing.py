"""Closing slide — three-theme recap cards."""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from . import theme as T


def render_closing(prs, *, title: str, themes: list[tuple[str, str, list[str]]],
                   footer: str = ""):
    """themes = [(theme_title, one_line_summary, [impact1, impact2, impact3]), ...]"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, T.SLIDE_W, T.SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = T.LIGHT_GRAY; bg.line.fill.background()

    # Title
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    T.apply_text(p.add_run(), title,
                 font_name=T.FONT_CN, size=T.FONT_SIZE["closing_title"],
                 bold=True, color=T.CRIMSON)

    # 3 cards horizontally
    n = min(len(themes), 3)
    if n == 0:
        return slide
    card_w = (T.SLIDE_W - Inches(1.2) - Inches(0.4) * (n - 1)) / n
    for i in range(n):
        left = Inches(0.6) + i * (card_w + Inches(0.4))
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                      left, Inches(1.7),
                                      card_w, Inches(4.8))
        card.fill.solid(); card.fill.fore_color.rgb = T.CRIMSON; card.line.fill.background()

        ttl, summary, impacts = themes[i]

        # Title inside the crimson card
        head_tb = slide.shapes.add_textbox(left + Inches(0.3), Inches(2.0),
                                           card_w - Inches(0.6), Inches(0.8))
        hp = head_tb.text_frame.paragraphs[0]
        T.apply_text(hp.add_run(), ttl,
                     font_name=T.FONT_CN, size=T.FONT_SIZE["closing_title"],
                     bold=True, color=T.WHITE)

        sum_tb = slide.shapes.add_textbox(left + Inches(0.3), Inches(2.9),
                                          card_w - Inches(0.6), Inches(1.0))
        sp = sum_tb.text_frame.paragraphs[0]
        T.apply_text(sp.add_run(), summary,
                     font_name=T.FONT_CN, size=T.FONT_SIZE["toc_item"],
                     color=T.LIGHT_GRAY)

        # Impacts list
        for j, imp in enumerate(impacts[:3]):
            im_tb = slide.shapes.add_textbox(left + Inches(0.3), Inches(4.2 + j * 0.55),
                                             card_w - Inches(0.6), Inches(0.55))
            ip = im_tb.text_frame.paragraphs[0]
            T.apply_text(ip.add_run(), f"· {imp}",
                         font_name=T.FONT_CN, size=T.FONT_SIZE["bullet"],
                         color=T.WHITE)

    if footer:
        ftb = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12), Inches(0.4))
        fp = ftb.text_frame.paragraphs[0]
        T.apply_text(fp.add_run(), footer,
                     font_name=T.FONT_EN, size=T.FONT_SIZE["pageno"],
                     color=T.MUTED_GRAY)
    return slide
